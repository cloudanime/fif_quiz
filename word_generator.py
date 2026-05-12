import os
import subprocess
import sys

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Installing python-docx...")
    install('python-docx')
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_manual_word(image_folder, output_name="app_user_manual.docx"):
    doc = Document()
    
    # Theme Colors
    PURPLE = RGBColor(69, 14, 78)
    YELLOW = RGBColor(255, 214, 0)
    
    # Title Page
    doc.add_paragraph("\n\n\n")
    title = doc.add_heading('Legacy Quiz', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = PURPLE
        run.font.size = Pt(48)
        
    subtitle = doc.add_paragraph('User Guide & Documentation')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(100, 100, 100)
        
    doc.add_paragraph("\n")
    motto = doc.add_paragraph("Celebrating 100 Years of God's Grace")
    motto.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in motto.runs:
        run.font.italic = True
        run.font.size = Pt(16)
        run.font.color.rgb = PURPLE

    doc.add_page_break()

    # List of images and their descriptions
    slides_data = [
        ("0 Instructions.png", "Guidance on playing different Game versions, which include:\n- Offline Single Player\n- Offline Multiple Players\n- Online Single Player\n- Online Individual Challenge\n- Online Team Challenge"),
        ("1 Legacy Continues.png", "Message from Visionary and Founder, Archbishop Professor Ezekiel H. Guti."),
        ("3 Offline Game.png", "Offline Game: Synchronize your local database by clicking any of the 2 buttons below. Once done you do not need internet connection to enjoy the game, just start the game."),
        ("4 Adding Players.png", "Renaming players is optional. Each game has a randomly selected 100 questions from our database. You determine how many questions will be played for your game session."),
        ("5 Offline Quiz Grid.png", "Quiz Grid: Scoreboard is located at the bottom of the questions grid. You can change the grid's width, height or font size by using the sliders above."),
        ("6 Mulitple Choice Questions.png", "Multiple Choice: Select the correct answer from the options provided. Immediate feedback will be shown in solo mode."),
        ("7 Marking Answers.png", "Marking Answers: Correct answer and source of the answer is highlighted for your further reading. If Managed by Quiz Master or if marking requires manual marking, use the grid below to mark participants' answers."),
        ("8 Picture Based Questions.png", "Picture Questions: Identify key historical figures, events, or artifacts. These can include book images from our official library."),
        ("9 Open or Structured Questions.png", "Structured Questions: Answer open-ended questions based on deep study of the Legacy materials."),
        ("questions grid.png", "Navigation: Questions that have been attempted will not be available for selections, the flip side of such questions will be shown on the grid. Questions are randomly selected from the grid by the player or Quiz Master."),
        ("11 Online (Quiz Master).png", "Online Hosting: The interface for the Quiz Master to manage live online sessions and track global progress."),
        ("12 Winner.png", "Victory: The celebratory reveal of the final standings and trophy distribution."),
        ("winner review.png", "Grand Review: A look at the final leaderboard and celebratory victory animation."),
        ("create team challenge.png", "Team Challenge: Create a group-based competition. This mode allows for large-scale event participation with multiple teams."),
        ("JoiningTeamChallenge.png", "Joining Teams: Players can enter their specific team code to join their designated group for the challenge."),
        ("team joining.png", "Team Lobby: Watch as your teammates join the session in real-time before the start."),
        ("13 Log in.png", "Authentication: Securely log in to access cloud-synced game features."),
        ("15 Register.png", "Registration: Join the Legacy Quiz community and start hosting your own sessions."),
    ]

    print("Generating Grand Word Manual...")

    for img_name, description in slides_data:
        img_path = os.path.join(image_folder, img_name)
        if not os.path.exists(img_path):
            print(f"Warning: {img_path} not found. Skipping.")
            continue
            
        # Clean up Title
        import re
        title_text = img_name.split('.')[0]
        title_text = re.sub(r'^\d+', '', title_text).replace('_', ' ').strip().upper()
        
        # Section Heading
        h = doc.add_heading(title_text, level=1)
        for run in h.runs:
            run.font.color.rgb = PURPLE
            
        # Layout: Table with 2 columns
        table = doc.add_table(rows=1, cols=2)
        table.autofit = True
        
        # Left Column: Image
        cell_img = table.rows[0].cells[0]
        paragraph = cell_img.paragraphs[0]
        run = paragraph.add_run()
        run.add_picture(img_path, width=Inches(2.5)) # Standard width for Word
        
        # Right Column: Description
        cell_desc = table.rows[0].cells[1]
        cell_desc.text = description
        
        doc.add_paragraph("\n") # Spacing

    # Footer
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "* For Guidance and not exhaustive."
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(output_name)
    print(f"Success! Your Grand Word manual is ready: {output_name}")

if __name__ == "__main__":
    image_dir = os.path.join("manual", "screenshots")
    if not os.path.exists(image_dir):
        print(f"Error: Folder '{image_dir}' not found.")
    else:
        create_manual_word(image_dir)
